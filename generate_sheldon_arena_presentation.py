#!/usr/bin/env python3
"""
Sheldon Arena - Self-Running Esports Club Presentation Generator
Generates a 16:9 widescreen PowerPoint (.pptx) with Sheldon School Purple & Gold branding.
Updated with:
- Slide transition duration increased by 50% (from 8s to 12s per slide).
- Terms 1 & 2 drop-in sessions: Open to all students individually / with friends (no team needed yet).
- Teams of 5 decided/formed by the end of Term 2 (Captain, Name, Tag, Logo, Intro Music).
- Mon-Fri Year Group Timetable.
- Year-Long Season Roadmap -> Grand Arena Finals at the end of the year.
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

# Sheldon School Branding Palette
COLOR_DEEP_PURPLE = RGBColor(38, 14, 66)      # #260E42 - Primary Deep Background
COLOR_DARK_PURPLE = RGBColor(52, 20, 90)      # #34145A - Secondary Background
COLOR_CARD_PURPLE = RGBColor(64, 25, 112)     # #401970 - Elevated Card Fill
COLOR_GOLD = RGBColor(255, 199, 44)           # #FFC72C - Sheldon School Primary Gold
COLOR_WARM_GOLD = RGBColor(245, 166, 35)      # #F5A623 - Amber Accent Gold
COLOR_LIGHT_GOLD = RGBColor(255, 230, 150)    # #FFE696 - Soft Highlight Gold
COLOR_WHITE = RGBColor(255, 255, 255)         # #FFFFFF - Clean Title Text
COLOR_SILVER = RGBColor(225, 220, 240)        # #E1DCF0 - Primary Body Text
COLOR_MUTED = RGBColor(180, 170, 205)         # #B4AACD - Secondary Subtext

SLIDE_DURATION_SECONDS = 12  # 50% longer than original 8s

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def set_slide_transition_timing(slide, advance_seconds=SLIDE_DURATION_SECONDS):
    """Configures automatic slide advance for self-running kiosk loops (12 seconds per slide)."""
    transition_xml = f"""
    <p:transition {nsdecls('p')} advTm="{advance_seconds * 1000}" advClick="1">
        <p:fade />
    </p:transition>
    """
    slide._element.append(parse_xml(transition_xml))

def add_header(slide, category, title, subtitle=None):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(13)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_GOLD

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.45), Inches(2.5), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_GOLD
    line.line.color.rgb = COLOR_GOLD

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(3.5), Inches(1.35), Inches(9.0), Inches(0.35))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = COLOR_MUTED

def add_card(slide, left, top, width, height, title, body_paragraphs, icon=None, badge=None, accent_gold=False):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_PURPLE
    card.line.color.rgb = COLOR_GOLD if accent_gold else RGBColor(180, 140, 50)
    card.line.width = Pt(2.0 if accent_gold else 1.2)

    pad = Inches(0.22)
    tb = slide.shapes.add_textbox(left + pad, top + pad, width - (pad * 2), height - (pad * 2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    if badge:
        p_b = tf.paragraphs[0]
        p_b.text = badge.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_GOLD
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]

    header_text = f"{icon} {title}" if icon else title
    p_title.text = header_text
    p_title.font.size = Pt(17)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_LIGHT_GOLD
    p_title.space_after = Pt(6)

    for item in body_paragraphs:
        p = tf.add_paragraph()
        if isinstance(item, tuple):
            prefix, text_content = item
            p.text = f"{prefix} {text_content}"
        else:
            p.text = f"• {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_SILVER
        p.space_after = Pt(4)

def build_presentation(output_path):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    script_dir = os.path.dirname(os.path.abspath(__file__))
    textures_dir = os.path.join(script_dir, "mods", "esports_core", "textures")

    # =========================================================================
    # SLIDE 1: Title Slide (Hero)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide1, SLIDE_DURATION_SECONDS)

    deco = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    deco.fill.solid()
    deco.fill.fore_color.rgb = COLOR_GOLD
    deco.line.fill.background()

    deco_bot = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.3), Inches(13.333), Inches(0.2))
    deco_bot.fill.solid()
    deco_bot.fill.fore_color.rgb = COLOR_GOLD
    deco_bot.line.fill.background()

    lion_logo = os.path.join(textures_dir, "esports_logo_lib_lion.png")
    dragon_logo = os.path.join(textures_dir, "esports_logo_lib_dragon.png")
    if os.path.exists(lion_logo):
        slide1.shapes.add_picture(lion_logo, Inches(1.0), Inches(2.0), width=Inches(2.5))
    if os.path.exists(dragon_logo):
        slide1.shapes.add_picture(dragon_logo, Inches(9.8), Inches(2.0), width=Inches(2.5))

    center_box = slide1.shapes.add_textbox(Inches(3.3), Inches(1.4), Inches(6.7), Inches(4.5))
    tf1 = center_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "★ OFFICIAL SHELDON SCHOOL ESPORTS ★"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_GOLD
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf1.add_paragraph()
    p1.text = "SHELDON ARENA"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(6)

    p2 = tf1.add_paragraph()
    p2.text = "COME & TRY IT • FORM YOUR SQUAD • LEAGUE MATCHES • GRAND FINALS"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WARM_GOLD
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)

    p3 = tf1.add_paragraph()
    p3.text = "Terms 1 & 2 drop-in sessions start now! No team needed yet — drop in on your year group's day, try the game, and form your 5-player squad by the end of Term 2."
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_SILVER
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(12)

    badge_shape = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.4), Inches(5.7), Inches(6.5), Inches(0.75)
    )
    badge_shape.fill.solid()
    badge_shape.fill.fore_color.rgb = COLOR_CARD_PURPLE
    badge_shape.line.color.rgb = COLOR_GOLD
    badge_shape.line.width = Pt(2.0)
    tf_badge = badge_shape.text_frame
    p_badge = tf_badge.paragraphs[0]
    p_badge.text = "⚡ NO SIGN-UP & NO TEAM NEEDED YET — JUST DROP IN! ⚡"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_GOLD
    p_badge.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 2: Weekly Schedule (Terms 1 & 2 Drop-Ins)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide2, SLIDE_DURATION_SECONDS)
    add_header(slide2, "Club Timetable", "Weekly Year-Group Schedule (Terms 1 & 2)", "Drop into the computing labs on your designated weekday — open to everyone!")

    d_w = Inches(2.25)
    d_h = Inches(4.9)
    d_y = Inches(1.75)

    add_card(
        slide2, Inches(0.8), d_y, d_w, d_h,
        "Monday",
        [
            "YEAR 7 SESSIONS",
            "Terms 1 & 2: Come & Try It open play.",
            "Drop in solo or with friends.",
            "Test weapons, movement, and cover building.",
            "Find your future squadmates!"
        ],
        icon="📅",
        badge="Year 7",
        accent_gold=True
    )

    add_card(
        slide2, Inches(3.2), d_y, d_w, d_h,
        "Tuesday",
        [
            "YEAR 8 SESSIONS",
            "Terms 1 & 2: Come & Try It open play.",
            "Casual drop-in matches.",
            "Try CTF, TDM & King of the Hill.",
            "Connect with other Year 8 players."
        ],
        icon="📅",
        badge="Year 8"
    )

    add_card(
        slide2, Inches(5.6), d_y, d_w, d_h,
        "Wednesday",
        [
            "YEAR 9 SESSIONS",
            "Terms 1 & 2: Come & Try It open play.",
            "Fast-paced tactical 3D combat.",
            "Experiment with squad setups.",
            "Get ready for team formation!"
        ],
        icon="📅",
        badge="Year 9"
    )

    add_card(
        slide2, Inches(8.0), d_y, d_w, d_h,
        "Thursday",
        [
            "YEAR 10 SESSIONS",
            "Terms 1 & 2: Come & Try It open play.",
            "High-intensity practice games.",
            "Master closing storm circle battles.",
            "Scout players for your 5-player team."
        ],
        icon="📅",
        badge="Year 10"
    )

    add_card(
        slide2, Inches(10.4), d_y, d_w, d_h,
        "Friday",
        [
            "YEAR 11 & 6TH FORM",
            "Senior drop-in gaming arena.",
            "Competitive open matches.",
            "Test tactics, ramps & high-ground cover.",
            "Build your championship lineup."
        ],
        icon="📅",
        badge="Y11 & 6th Form",
        accent_gold=True
    )

    # =========================================================================
    # SLIDE 3: Build Your Squad of 5 (Formed by End of Term 2)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide3, SLIDE_DURATION_SECONDS)
    add_header(slide3, "Team Formation", "Form Your Team of 5 (By End of Term 2)", "During Terms 1 & 2, play openly — then register your official 5-player team!")

    add_card(
        slide3, Inches(0.8), Inches(1.75), Inches(3.7), Inches(4.9),
        "1. Full 5-Player Team",
        [
            "Decided by the end of Term 2.",
            "Must be a complete team of 5 players from your year group.",
            "All students must be part of a team for league fixtures.",
            "Appoint 1 Team Captain to lead squad comms & match lineups."
        ],
        icon="🛡️",
        badge="Squad Lineup",
        accent_gold=True
    )

    add_card(
        slide3, Inches(4.8), Inches(1.75), Inches(3.7), Inches(4.9),
        "2. Team Name & 3-Letter Tag",
        [
            "Create a unique squad name (e.g. Chippenham Knights, Apex Titans).",
            "Pick an official 3-letter in-game Tag (e.g. [KNT], [APX], [SHD]).",
            "Your tag is displayed in-game on nametags, HUD, and live match killfeeds!"
        ],
        icon="🏷️",
        badge="Team Brand"
    )

    add_card(
        slide3, Inches(8.8), Inches(1.75), Inches(3.7), Inches(4.9),
        "3. Logo & Walk-In Music",
        [
            "🎨 Team Logo: Pick your squad mascot (Lion, Dragon, Eagle, etc.) for live scoreboards & broadcasts.",
            "🎵 Custom Intro Music: Select your squad's signature hype walk-in music that plays dynamically on arena entry!"
        ],
        icon="🎵",
        badge="Theme & Insignia",
        accent_gold=True
    )

    # =========================================================================
    # SLIDE 4: Game Modes
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide4, SLIDE_DURATION_SECONDS)
    add_header(slide4, "Match Types", "Arena Game Modes", "Four intense competitive battle formats played throughout the season.")

    g_w = Inches(5.7)
    g_h = Inches(2.4)

    add_card(
        slide4, Inches(0.8), Inches(1.75), g_w, g_h,
        "Capture The Flag (CTF)",
        [
            "Infiltrate the enemy base, grab their flag, and sprint home.",
            "Carrier lockout: Flag carriers cannot shoot — squads must provide suppressive cover fire!"
        ],
        icon="🚩",
        badge="5v5 Team Strategy"
    )

    add_card(
        slide4, Inches(6.8), Inches(1.75), g_w, g_h,
        "Team Deathmatch (TDM)",
        [
            "Squad elimination shootout with tactical building and weapon looting.",
            "A shrinking electric storm perimeter forces climactic final showdowns."
        ],
        icon="⚔️",
        badge="High Intensity"
    )

    add_card(
        slide4, Inches(0.8), Inches(4.45), g_w, g_h,
        "King of the Hill (KOTH)",
        [
            "Dominate a shifting central capture hill against heavy enemy assault.",
            "Hold the zone under fire and deploy barricades to secure points per second."
        ],
        icon="👑",
        badge="Zone Control"
    )

    add_card(
        slide4, Inches(6.8), Inches(4.45), g_w, g_h,
        "Payload & Domination",
        [
            "Payload: Escort a moving cart into enemy territory while defenders contest.",
            "Domination: Coordinate squad splits to capture and maintain 3 key control points."
        ],
        icon="🚚",
        badge="Objective Assault"
    )

    # =========================================================================
    # SLIDE 5: The Year-Long Season Roadmap
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide5, SLIDE_DURATION_SECONDS)
    add_header(slide5, "Season Timeline", "The Sheldon Arena Season Roadmap", "How the club progresses from first drop-ins to the championship finale.")

    p_w = Inches(2.7)
    p_h = Inches(4.9)
    p_y = Inches(1.75)

    add_card(
        slide5, Inches(0.8), p_y, p_w, p_h,
        "Phase 1: Terms 1 & 2",
        [
            "\"Come & Try It\" Sessions.",
            "Drop in on your Year Group's day.",
            "NO TEAMS NEEDED YET!",
            "Play, test weapons & building.",
            "Meet players and form squads."
        ],
        icon="🎯",
        badge="Terms 1 & 2"
    )

    add_card(
        slide5, Inches(3.8), p_y, p_w, p_h,
        "Phase 2: End of Term 2",
        [
            "Official Team Registration.",
            "Finalize your 5-player squad.",
            "Pick captain, tag, logo & music.",
            "Dedicated team practice times.",
            "Warm up against AI Sentry Bots."
        ],
        icon="🛡️",
        badge="Team Practices"
    )

    add_card(
        slide5, Inches(6.8), p_y, p_w, p_h,
        "Phase 3: League Matches",
        [
            "Official scheduled match fixtures.",
            "Weekly competitive rounds.",
            "Earn league points and win/loss records.",
            "Climb the live school standings ladder."
        ],
        icon="📊",
        badge="Season League"
    )

    add_card(
        slide5, Inches(9.8), p_y, p_w, p_h,
        "Phase 4: Arena Finals",
        [
            "End-of-Year Championship!",
            "The top teams compete in an arena-style live event.",
            "Spectators, live casting & big screens.",
            "Championship trophies & school glory!"
        ],
        icon="🏆",
        badge="Grand Finale",
        accent_gold=True
    )

    # =========================================================================
    # SLIDE 6: Tactical Combat Mechanics
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide6, SLIDE_DURATION_SECONDS)
    add_header(slide6, "Gameplay Mechanics", "Tactical Combat & Arena Features", "Key competitive mechanics your squad will master during matches.")

    add_card(
        slide6, Inches(0.8), Inches(1.75), g_w, g_h,
        "Instant Tactical Building",
        [
            "Deploy wooden barricades, walls, and ramps in real-time under fire.",
            "Build instant cover, block sniper sightlines, or rush high ground advantage."
        ],
        icon="🧱",
        badge="Real-Time Cover"
    )

    add_card(
        slide6, Inches(6.8), Inches(1.75), g_w, g_h,
        "Electric Closing Storm",
        [
            "A contracting electric storm perimeter steadily forces teams to the island center.",
            "Guarantees intense, action-packed final-circle finishes with zero camping."
        ],
        icon="⚡",
        badge="Dynamic Perimeter"
    )

    add_card(
        slide6, Inches(0.8), Inches(4.45), g_w, g_h,
        "PvE Sentry Bot Warmups",
        [
            "Dedicated practice range featuring AI Sentry Bots across 3 difficulty tiers.",
            "Great for warming up your aim and testing strategies before league matches."
        ],
        icon="🤖",
        badge="Practice Range"
    )

    add_card(
        slide6, Inches(6.8), Inches(4.45), g_w, g_h,
        "Live In-Game Stats & Ratings",
        [
            "Dynamic scoreboard tracks kills, deaths, captures, and player ratings.",
            "Live broadcast spectator cameras and automatic killfeed overlays."
        ],
        icon="📈",
        badge="Performance Rating"
    )

    # =========================================================================
    # SLIDE 7: How to Get Involved / Next Steps
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, COLOR_DEEP_PURPLE)
    set_slide_transition_timing(slide7, SLIDE_DURATION_SECONDS)
    add_header(slide7, "Get Ready", "How to Enter Sheldon Arena", "No sign-up and no team needed yet — just drop in on your year group's day!")

    cta_card = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.75), Inches(10.9), Inches(4.9)
    )
    cta_card.fill.solid()
    cta_card.fill.fore_color.rgb = COLOR_CARD_PURPLE
    cta_card.line.color.rgb = COLOR_GOLD
    cta_card.line.width = Pt(2.0)

    tf_c = cta_card.text_frame
    tf_c.word_wrap = True

    p_c0 = tf_c.paragraphs[0]
    p_c0.text = "HOW TO GET STARTED (3 EASY STEPS):"
    p_c0.font.size = Pt(20)
    p_c0.font.bold = True
    p_c0.font.color.rgb = COLOR_GOLD
    p_c0.space_after = Pt(12)

    steps = [
        ("1️⃣", "Drop In on Your Year Group's Day (Terms 1 & 2)", "Mon: Y7 • Tue: Y8 • Wed: Y9 • Thu: Y10 • Fri: Y11 & Sixth Form. No team or sign-up needed!"),
        ("2️⃣", "Play, Test Weapons & Form Your Squad", "Try the game in Terms 1 & 2, meet other players, and assemble your 5-player team by the end of Term 2."),
        ("3️⃣", "Register Your Team for the League", "Choose your Team Captain, Team Name, 3-letter Tag, Mascot Logo, and signature Walk-In Music for the league!")
    ]

    for num, title, desc in steps:
        p_step = tf_c.add_paragraph()
        p_step.text = f"{num}  {title} — {desc}"
        p_step.font.size = Pt(13.5)
        p_step.font.color.rgb = COLOR_WHITE
        p_step.space_after = Pt(8)

    p_foot = tf_c.add_paragraph()
    p_foot.text = "🏆 TERMS 1 & 2: COME & TRY IT • TEAMS FORMED BY END OF TERM 2 • GRAND FINALS AT END OF YEAR! 🏆"
    p_foot.font.size = Pt(13.5)
    p_foot.font.bold = True
    p_foot.font.color.rgb = COLOR_LIGHT_GOLD
    p_foot.space_before = Pt(8)

    prs.save(output_path)
    print(f"Updated presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Sheldon_Arena_Presentation.pptx")
    build_presentation(out_file)
